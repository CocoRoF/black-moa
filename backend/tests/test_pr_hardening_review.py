"""Regressions for the security-hardening review (credits, retention, SSRF, parser, keys).

Each test here covers a defect found while auditing the hardening PR against
plan/15 and plan/19, or an unproven claim those plans make.
"""
from __future__ import annotations

import asyncio
import io
import subprocess
import sys
import uuid
from datetime import UTC, datetime, timedelta
from types import SimpleNamespace

import pytest
from sqlalchemy import func, select, text

from memora.core.errors import PaymentRequired
from memora.db.session import SessionLocal, session_scope
from memora.models import (
    Agent,
    Conversation,
    CreditLedger,
    CreditReservation,
    Job,
    Turn,
    UsageEvent,
    User,
    Visitor,
)
from memora.services import credits as CR
from tests.conftest import auth, signup
from tests.test_domain import _run_jobs

# asyncio_mode = "auto" collects the async cases; the sync ones stay unmarked so
# they do not emit "marked with asyncio but not async" warnings.


async def _fund(owner_id: uuid.UUID, amount: str) -> None:
    """Set the owner's balance to exactly ``amount`` credits."""
    async with session_scope() as db:
        current = await CR.balance(db, owner_id)
        delta = CR.q(amount) - current
        if delta:
            await CR.apply(db, owner_id, delta, "adjust", note="test funding", allow_reserved=True)


async def _agent_and_conversation(owner_id: uuid.UUID) -> tuple[uuid.UUID, uuid.UUID]:
    async with session_scope() as db:
        agent = Agent(owner_id=owner_id, name="Reaper Test", provider="fake", model_id="fake-1")
        db.add(agent)
        await db.flush()
        conv = Conversation(owner_id=owner_id, agent_id=agent.id, audience="owner")
        db.add(conv)
        await db.flush()
        return agent.id, conv.id


# ── credits: recorded usage must never be refused ────────────────────────────


async def test_charge_usage_is_recorded_while_a_turn_holds_credits(client):
    """STT/TTS/embedding cost is already incurred; refusing it desynchronises the ledger."""
    user, _ = await signup(client)
    owner_id = uuid.UUID(user["id"])
    await _fund(owner_id, "5")
    turn_id = uuid.uuid4()
    async with session_scope() as db:
        await CR.reserve_turn(db, owner_id=owner_id, turn_id=turn_id, turn_cap=5, daily_cap=None, audience="owner")

    async with session_scope() as db:
        _, reserved, available = await CR.balance_state(db, owner_id)
        assert (float(reserved), float(available)) == (5.0, 0.0)
        await CR.charge_usage(db, owner_id=owner_id, kind="stt", credits="1.5", provider="openai", units=0.5)

    async with session_scope() as db:
        total, reserved, _ = await CR.balance_state(db, owner_id)
        assert float(total) == 3.5
        assert float(reserved) == 5.0  # the hold itself is untouched
        stt = (await db.execute(select(UsageEvent).where(UsageEvent.owner_id == owner_id,
                                                          UsageEvent.kind == "stt"))).scalars().all()
        assert [float(e.credits) for e in stt] == [1.5]
        # a discretionary (admin) debit is still refused while credits are held
        with pytest.raises(PaymentRequired) as e:
            await CR.apply(db, owner_id, -1, "adjust", note="admin clawback")
        assert e.value.code == "credits_reserved"


# ── credits: leaked holds must expire ────────────────────────────────────────


async def test_stale_reservation_is_released_by_the_reaper(client):
    """A hard-killed process must not shrink the owner's balance forever."""
    user, _ = await signup(client)
    owner_id = uuid.UUID(user["id"])
    await _fund(owner_id, "10")
    agent_id, conv_id = await _agent_and_conversation(owner_id)

    async with session_scope() as db:
        turn = Turn(conversation_id=conv_id, owner_id=owner_id, agent_id=agent_id, audience="owner",
                    status="running", started_at=datetime.now(UTC) - timedelta(hours=6))
        db.add(turn)
        await db.flush()
        turn_id = turn.id
        await CR.reserve_turn(db, owner_id=owner_id, turn_id=turn_id, turn_cap=10, daily_cap=None, audience="owner")
        res = await db.get(CreditReservation, turn_id)
        res.created_at = datetime.now(UTC) - timedelta(hours=6)

    async with session_scope() as db:
        assert float(await CR.available_balance(db, owner_id)) == 0.0
        # a young hold is left alone
        assert await CR.expire_stale_reservations(db, minutes=60 * 24) == 0

    async with session_scope() as db:
        assert await CR.expire_stale_reservations(db, minutes=120) >= 1

    async with session_scope() as db:
        assert float(await CR.available_balance(db, owner_id)) == 10.0
        assert (await db.get(CreditReservation, turn_id)).status == "released"
        healed = await db.get(Turn, turn_id)
        assert healed.status == "failed" and healed.error_code == "turn_abandoned"
        # idempotent
        assert await CR.expire_stale_reservations(db, minutes=120) == 0


async def test_reservation_reaper_is_registered_and_scheduled():
    from memora.worker.__main__ import SCHEDULE
    from memora.worker.handlers import HANDLERS

    assert "credits.reservation_reaper" in HANDLERS
    assert "credits.reservation_reaper" in {kind for kind, _every, _dedupe in SCHEDULE}
    # the retention boundary registered in handlers.py must be the full one
    assert HANDLERS["retention.sweep"].__module__.endswith("worker.handlers")
    assert "retention.purge_memory" in HANDLERS


async def test_reservation_reaper_handler_sweeps_terminal_audit_rows(client):
    user, _ = await signup(client)
    owner_id = uuid.UUID(user["id"])
    old = uuid.uuid4()
    async with session_scope() as db:
        db.add(CreditReservation(turn_id=old, owner_id=owner_id, amount=1, day=await CR.usage_day(db, owner_id),
                                 audience="owner", status="settled", created_at=datetime.now(UTC) - timedelta(days=90),
                                 settled_at=datetime.now(UTC) - timedelta(days=90)))
    async with session_scope() as db:
        from memora.worker.handlers import HANDLERS

        result = await HANDLERS["credits.reservation_reaper"](db, {})
    assert result["swept"] >= 1
    async with session_scope() as db:
        assert await db.get(CreditReservation, old) is None


# ── credits: real concurrency ────────────────────────────────────────────────


async def test_two_concurrent_turns_cannot_both_pass_the_cap(client):
    """Two overlapping sessions must serialise on the CreditBalance row lock."""
    user, _ = await signup(client)
    owner_id = uuid.UUID(user["id"])
    await _fund(owner_id, "5")
    t1, t2 = uuid.uuid4(), uuid.uuid4()

    async def reserve(turn_id):
        async with SessionLocal() as db:
            try:
                res = await CR.reserve_turn(db, owner_id=owner_id, turn_id=turn_id, turn_cap=5,
                                            daily_cap=None, audience="owner")
                amount = float(res.amount)
                await asyncio.sleep(0.05)  # hold the row lock across the other task's attempt
                await db.commit()
                return amount
            except PaymentRequired as e:
                await db.rollback()
                return e

    a, b = await asyncio.gather(reserve(t1), reserve(t2))
    outcomes = [a, b]
    granted = [x for x in outcomes if isinstance(x, float)]
    refused = [x for x in outcomes if isinstance(x, PaymentRequired)]
    assert len(granted) == 1 and granted[0] == 5.0, outcomes
    assert len(refused) == 1

    async with session_scope() as db:
        total, reserved, available = await CR.balance_state(db, owner_id)
        assert float(reserved) <= float(total)
        assert float(available) == 0.0
        held = await db.scalar(select(func.count()).select_from(CreditReservation)
                               .where(CreditReservation.owner_id == owner_id, CreditReservation.status == "held"))
        assert held == 1


async def test_concurrent_visitor_turns_respect_the_daily_turn_cap(client):
    user, _ = await signup(client)
    owner_id = uuid.UUID(user["id"])
    await _fund(owner_id, "50")

    async def reserve(turn_id):
        async with SessionLocal() as db:
            try:
                await CR.reserve_turn(db, owner_id=owner_id, turn_id=turn_id, turn_cap=5, daily_cap=None,
                                      audience="visitor", visitor_turn_cap=1)
                await asyncio.sleep(0.05)
                await db.commit()
                return True
            except Exception:
                await db.rollback()
                return False

    results = await asyncio.gather(*(reserve(uuid.uuid4()) for _ in range(4)))
    assert sum(results) == 1, results


# ── credits: settle invariants ───────────────────────────────────────────────


async def test_settle_over_reservation_keeps_ledger_equal_to_usage_events(client):
    """Provider overage is capped, audited on the reservation, never charged twice."""
    user, _ = await signup(client)
    owner_id = uuid.UUID(user["id"])
    await _fund(owner_id, "4")
    turn_id, agent_id = uuid.uuid4(), uuid.uuid4()
    async with session_scope() as db:
        await CR.reserve_turn(db, owner_id=owner_id, turn_id=turn_id, turn_cap=3, daily_cap=None, audience="owner")

    async with session_scope() as db:
        available, charged = await CR.settle_turn(
            db, owner_id=owner_id, agent_id=agent_id, turn_id=turn_id, provider="fake", model_id="fake-1",
            input_tokens=1_000_000, output_tokens=1_000_000, cache_read=0, cache_write=0, cost_usd=9.0,
            credits="12.5", audience="owner")
        assert float(charged) == 3.0  # capped at the hold, not 12.5
        assert float(available) == 1.0

    async with session_scope() as db:
        res = await db.get(CreditReservation, turn_id)
        assert res.status == "settled"
        assert float(res.actual_credits) == 12.5 and float(res.charged_credits) == 3.0
        total, reserved, _ = await CR.balance_state(db, owner_id)
        assert (float(total), float(reserved)) == (1.0, 0.0)
        ledger = await db.scalar(select(func.coalesce(func.sum(CreditLedger.delta), 0))
                                 .where(CreditLedger.owner_id == owner_id, CreditLedger.kind == "turn"))
        usage = await db.scalar(select(func.coalesce(func.sum(UsageEvent.credits), 0))
                                .where(UsageEvent.owner_id == owner_id, UsageEvent.kind == "llm"))
        assert float(-ledger) == float(usage) == 3.0
        # replay is idempotent
        _, again = await CR.settle_turn(
            db, owner_id=owner_id, agent_id=agent_id, turn_id=turn_id, provider="fake", model_id="fake-1",
            input_tokens=1, output_tokens=1, cache_read=0, cache_write=0, cost_usd=0, credits="12.5", audience="owner")
        assert float(again) == 3.0
        assert await db.scalar(select(func.count()).select_from(CreditLedger)
                               .where(CreditLedger.ref_type == "turn", CreditLedger.ref_id == str(turn_id))) == 1


async def test_settlement_is_booked_on_the_reservation_day_not_todays_row(client):
    """A turn started at 23:59 owner-local settles into the day it reserved."""
    user, _ = await signup(client)
    owner_id = uuid.UUID(user["id"])
    await _fund(owner_id, "20")
    async with session_scope() as db:
        owner = await db.get(User, owner_id)
        owner.timezone = "Pacific/Kiritimati"  # UTC+14: local date is ahead of UTC
    turn_id = uuid.uuid4()
    async with session_scope() as db:
        reserve_day = await CR.usage_day(db, owner_id)
        res = await CR.reserve_turn(db, owner_id=owner_id, turn_id=turn_id, turn_cap=5, daily_cap=100,
                                    audience="owner")
        assert res.day == reserve_day

    async with session_scope() as db:
        # simulate the local day rolling over before the turn finishes
        res = await db.get(CreditReservation, turn_id)
        res.day = reserve_day - timedelta(days=1)
        booked_day = res.day
    async with session_scope() as db:
        await CR.settle_turn(db, owner_id=owner_id, agent_id=uuid.uuid4(), turn_id=turn_id, provider="fake",
                             model_id="fake-1", input_tokens=10, output_tokens=10, cache_read=0, cache_write=0,
                             cost_usd=0.0, credits="2", audience="owner")
    async with session_scope() as db:
        row = (await db.execute(text("SELECT credits, reserved_credits, turns FROM usage_daily "
                                     "WHERE owner_id=:o AND day=:d"),
                                {"o": owner_id, "d": booked_day})).first()
        assert row is not None and float(row[0]) == 2.0
        assert float(row[1]) == 0.0  # the hold was released from the same row it was taken from


# ── retention ────────────────────────────────────────────────────────────────


async def test_retention_keeps_a_visitor_that_still_has_a_retained_conversation(client):
    from memora.services.retention import retention_sweep

    user, _ = await signup(client)
    owner_id = uuid.UUID(user["id"])
    old = datetime.now(UTC) - timedelta(days=200)
    async with session_scope() as db:
        agent = Agent(owner_id=owner_id, name="Partial Retention", provider="fake", model_id="fake-1",
                      visitor_settings={"retention_days": 30})
        db.add(agent)
        await db.flush()
        visitor = Visitor(owner_id=owner_id, agent_id=agent.id, token_hash=f"keep-{uuid.uuid4().hex}",
                          display_name="Still Active", email="keep@example.com", first_seen_at=old, last_seen_at=old)
        db.add(visitor)
        await db.flush()
        expired = Conversation(owner_id=owner_id, agent_id=agent.id, audience="visitor", visitor_id=visitor.id,
                               title="expired", last_message_at=old)
        fresh = Conversation(owner_id=owner_id, agent_id=agent.id, audience="visitor", visitor_id=visitor.id,
                             title="fresh", last_message_at=datetime.now(UTC))
        db.add_all([expired, fresh])
        await db.flush()
        visitor_id, expired_id, fresh_id = visitor.id, expired.id, fresh.id

    async with session_scope() as db:
        await retention_sweep(db, {})

    async with session_scope() as db:
        assert await db.get(Conversation, expired_id) is None
        assert await db.get(Conversation, fresh_id) is not None
        kept = await db.get(Visitor, visitor_id)
        assert kept is not None and kept.email == "keep@example.com"


async def test_retention_batches_job_payload_purge_and_keeps_the_memory_job(client):
    from memora.services import jobs as J
    from memora.services.retention import retention_sweep

    user, _ = await signup(client)
    owner_id = uuid.UUID(user["id"])
    old = datetime.now(UTC) - timedelta(days=200)
    async with session_scope() as db:
        agent = Agent(owner_id=owner_id, name="Job Purge", provider="fake", model_id="fake-1",
                      visitor_settings={"retention_days": 30})
        db.add(agent)
        await db.flush()
        visitor = Visitor(owner_id=owner_id, agent_id=agent.id, token_hash=f"jobs-{uuid.uuid4().hex}",
                          display_name="Gone", first_seen_at=old, last_seen_at=old)
        db.add(visitor)
        await db.flush()
        visitor_id = visitor.id
        await J.enqueue(db, "notify.evaluate", {"visitor_id": str(visitor_id), "text": "PII copy"})

    async with session_scope() as db:
        result = await retention_sweep(db, {})
        assert result["memory_purges_queued"] >= 1

    async with session_scope() as db:
        assert await db.get(Visitor, visitor_id) is None
        leftovers = (await db.execute(select(Job).where(Job.payload["visitor_id"].astext == str(visitor_id)))).scalars().all()
        assert [j.kind for j in leftovers] == ["retention.purge_memory"]
        # second pass changes nothing
        again = await retention_sweep(db, {})
        assert again["visitors_deleted"] == 0


async def test_retention_drops_orphaned_visitor_private_facts(client):
    from memora.models import Fact
    from memora.services.retention import retention_sweep

    user, _ = await signup(client)
    owner_id = uuid.UUID(user["id"])
    async with session_scope() as db:
        agent = Agent(owner_id=owner_id, name="Fact Retention", provider="fake", model_id="fake-1")
        db.add(agent)
        await db.flush()
        # visitor row is recent (kept), but every thread it appeared in is gone
        visitor = Visitor(owner_id=owner_id, agent_id=agent.id, token_hash=f"facts-{uuid.uuid4().hex}",
                          first_seen_at=datetime.now(UTC), last_seen_at=datetime.now(UTC))
        db.add(visitor)
        await db.flush()
        fact = Fact(owner_id=owner_id, agent_id=agent.id, visitor_id=visitor.id, subject="visitor",
                    predicate="phone", object="010-1234-5678", visibility="visitor_private")
        db.add(fact)
        await db.flush()
        fact_id, visitor_id = fact.id, visitor.id

    async with session_scope() as db:
        await retention_sweep(db, {})
    async with session_scope() as db:
        assert await db.get(Fact, fact_id) is None
        assert await db.get(Visitor, visitor_id) is not None


# ── SSRF ─────────────────────────────────────────────────────────────────────


def _stub_connection(monkeypatch, response: bytes, captured: dict):
    class Writer:
        def __init__(self):
            self.data = bytearray()

        def write(self, data):
            self.data.extend(data)

        async def drain(self):
            return None

        def close(self):
            return None

        async def wait_closed(self):
            return None

    writer = Writer()

    async def fake_open_connection(host, port, **kwargs):
        captured.setdefault("hosts", []).append(host)
        captured["kwargs"] = kwargs
        captured["writer"] = writer
        reader = asyncio.StreamReader()
        reader.feed_data(response)
        reader.feed_eof()
        return reader, writer

    from memora.services import safe_http as SH

    monkeypatch.setattr(SH.asyncio, "open_connection", fake_open_connection)


def test_safe_http_rejects_literal_internal_addresses(monkeypatch):
    from memora.services import safe_http as SH

    monkeypatch.setattr(SH, "get_settings", lambda: SimpleNamespace(outbound_allowed_ports="80,443"))
    for url in ("http://169.254.169.254/latest/meta-data/", "http://127.0.0.1/", "http://[::1]/",
                "http://[::ffff:127.0.0.1]/", "http://2130706433/", "http://0177.0.0.1/", "http://10.0.0.1/",
                "http://[fd00::1]/"):
        with pytest.raises(ValueError, match="blocked_host"):
            SH.resolve_target(url)
    with pytest.raises(ValueError, match="blocked_port"):
        SH.resolve_target("http://example.com:8080/")
    with pytest.raises(ValueError, match="url_credentials_forbidden"):
        SH.resolve_target("http://user:pw@example.com/")
    with pytest.raises(ValueError, match="invalid_url"):
        SH.resolve_target("file:///etc/passwd")


def test_outbound_port_allowlist_parsing(monkeypatch):
    from memora.services import safe_http as SH

    def ports(raw):
        monkeypatch.setattr(SH, "get_settings", lambda: SimpleNamespace(outbound_allowed_ports=raw))
        return sorted(SH._allowed_ports())

    assert ports("80,443") == [80, 443]
    assert ports(" 80 , 443 ,") == [80, 443]
    assert ports("") == [80, 443]  # empty config falls back to the documented default
    for bad in ("abc", "0", "70000", "80;443"):
        monkeypatch.setattr(SH, "get_settings", lambda bad=bad: SimpleNamespace(outbound_allowed_ports=bad))
        with pytest.raises(ValueError, match="invalid_outbound_port_config"):
            SH._allowed_ports()


async def test_safe_http_https_keeps_hostname_for_sni_while_pinning_the_ip(monkeypatch):
    from memora.services import safe_http as SH

    monkeypatch.setattr(SH, "_public_ips", lambda host, port: ("93.184.216.34",))
    monkeypatch.setattr(SH, "get_settings", lambda: SimpleNamespace(outbound_allowed_ports="80,443"))
    captured: dict = {}
    _stub_connection(monkeypatch, b"HTTP/1.1 200 OK\r\nContent-Length: 2\r\n\r\nok", captured)
    r = await SH.safe_request("GET", "https://example.com/x")
    assert r.status_code == 200
    assert captured["hosts"] == ["93.184.216.34"]
    assert captured["kwargs"]["server_hostname"] == "example.com"
    assert captured["kwargs"]["ssl"] is not None
    assert b"Host: example.com\r\n" in captured["writer"].data


async def test_safe_http_revalidates_every_redirect_hop(monkeypatch):
    from memora.services import safe_http as SH

    monkeypatch.setattr(SH, "get_settings", lambda: SimpleNamespace(outbound_allowed_ports="80,443"))

    def fake_resolve(url):
        if "evil.example" in url:
            raise ValueError("blocked_host")  # what the real resolver does for a private answer
        return SH._Target(url=url, scheme="http", host="public.example", port=80, request_target="/",
                          host_header="public.example", ips=("93.184.216.34",))

    monkeypatch.setattr(SH, "resolve_target", fake_resolve)
    captured: dict = {}
    _stub_connection(monkeypatch,
                     b"HTTP/1.1 302 Found\r\nLocation: http://evil.example/\r\nContent-Length: 0\r\n\r\n",
                     captured)
    with pytest.raises(ValueError, match="blocked_host"):
        await SH.safe_request("GET", "http://public.example/", max_redirects=3)


async def test_webhook_delivery_goes_through_safe_http(monkeypatch, client):
    from memora.services import notifications as NT

    calls: dict = {}

    async def fake_safe_request(method, url, **kwargs):
        calls.update(method=method, url=url, **kwargs)
        return SimpleNamespace(status_code=200, headers={}, content=b"", url=url)

    monkeypatch.setattr(NT, "safe_request", fake_safe_request)
    monkeypatch.setattr(NT, "check_url", lambda url: url)
    async with session_scope() as db:
        await NT.send_via(db, "webhook", {"url": "https://hooks.example/x", "secret": "s3cret"},
                          subject="s", text="t", html="<p>t</p>", payload={"event": "test"})
    assert calls["method"] == "POST" and calls["url"] == "https://hooks.example/x"
    assert calls["max_redirects"] == 0
    assert calls["headers"]["X-Memora-Signature"].startswith("sha256=")

    # a destination that stopped resolving publicly is refused at delivery time
    def blocked(_url):
        raise ValueError("blocked_host")

    monkeypatch.setattr(NT, "check_url", blocked)
    from memora.core.errors import ValidationFailed

    async with session_scope() as db:
        with pytest.raises(ValidationFailed):
            await NT.send_via(db, "slack", {"url": "https://hooks.example/x"}, subject="s", text="t", html="",
                              payload={})


# ── parser isolation ─────────────────────────────────────────────────────────


def test_parser_child_applies_rlimits_and_blocks_sockets():
    from memora.services import extract as EX

    probe = (
        "from memora.services.extract_worker import _apply_limits, _disable_network\n"
        "_apply_limits(); _disable_network()\n"
        "import json, resource, socket\n"
        "out = {'as': resource.getrlimit(resource.RLIMIT_AS)[0],\n"
        "       'cpu': resource.getrlimit(resource.RLIMIT_CPU)[0],\n"
        "       'nofile': resource.getrlimit(resource.RLIMIT_NOFILE)[0],\n"
        "       'fsize': resource.getrlimit(resource.RLIMIT_FSIZE)[0]}\n"
        "for name, fn in (('connect', lambda: socket.socket().connect(('93.184.216.34', 80))),\n"
        "                 ('create_connection', lambda: socket.create_connection(('1.1.1.1', 80))),\n"
        "                 ('getaddrinfo', lambda: socket.getaddrinfo('example.com', 80))):\n"
        "    try:\n"
        "        fn(); out[name] = 'ALLOWED'\n"
        "    except PermissionError as e:\n"
        "        out[name] = str(e)\n"
        "print(json.dumps(out))\n"
    )
    proc = subprocess.run([sys.executable, "-c", probe], capture_output=True, env=EX._sandbox_env(), text=True)
    assert proc.returncode == 0, proc.stderr
    import json as _json

    out = _json.loads(proc.stdout)
    assert out["as"] == 512 * 1024 * 1024
    assert out["cpu"] == 20 and out["nofile"] == 64 and out["fsize"] == 16 * 1024 * 1024
    assert out["connect"] == out["create_connection"] == out["getaddrinfo"] == "parser_network_disabled"


def test_real_parser_child_extracts_every_supported_office_format():
    import docx
    import openpyxl
    from pptx import Presentation

    from memora.services.extract import extract

    assert "hello world" in extract(b"# Title\n\nhello world", "text/markdown", "a.md").text

    buf = io.BytesIO()
    d = docx.Document()
    d.add_heading("Doc Heading", 1)
    d.add_paragraph("docx body line")
    d.save(buf)
    docx_text = extract(buf.getvalue(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                        "a.docx").text
    assert "# Doc Heading" in docx_text and "docx body line" in docx_text

    buf = io.BytesIO()
    prs = Presentation()
    prs.slides.add_slide(prs.slide_layouts[5]).shapes.title.text = "Slide Title"
    prs.save(buf)
    assert "Slide Title" in extract(
        buf.getvalue(), "application/vnd.openxmlformats-officedocument.presentationml.presentation", "a.pptx").text

    buf = io.BytesIO()
    wb = openpyxl.Workbook()
    wb.active["A1"], wb.active["B1"] = "cellA", "cellB"
    wb.save(buf)
    assert "cellA | cellB" in extract(
        buf.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "a.xlsx").text

    assert "isolated" in extract(b"<html><body><p>isolated parser</p></body></html>", "text/html", "a.html").text


def test_parser_child_rejects_a_zip_bomb_without_killing_the_parent():
    import zipfile

    from memora.services.extract import extract

    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zf:
        zf.writestr("[Content_Types].xml", "<Types/>")
        zf.writestr("word/document.xml", b"0" * (64 * 1024 * 1024))
    with pytest.raises(ValueError) as e:
        extract(buf.getvalue(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document", "bomb.docx")
    assert "archive_" in str(e.value)


async def test_parser_timeout_produces_a_failed_document_not_a_stuck_job(client, monkeypatch):
    from memora.services import extract as EX
    from memora.services import knowledge as K

    def timeout(*_args, **kwargs):
        raise subprocess.TimeoutExpired(cmd="parser", timeout=kwargs.get("timeout", 1))

    monkeypatch.setattr(EX.subprocess, "run", timeout)
    with pytest.raises(ValueError, match="parser_timeout"):
        EX.extract(b"x", "application/pdf", "slow.pdf")

    user, tok = await signup(client)
    files = {"file": ("slow.pdf", b"%PDF-1.4 broken", "application/pdf")}
    doc = (await client.post("/api/knowledge/documents", files=files, data={"kind": "file"}, headers=auth(tok))).json()
    async with session_scope() as db:
        out = await K.index_document(db, uuid.UUID(doc["id"]))
    assert out.get("failed") == "parser_timeout"
    got = (await client.get(f"/api/knowledge/documents/{doc['id']}", headers=auth(tok))).json()
    assert got["status"] == "failed" and got["error"] == "parser_timeout"


async def test_knowledge_pipeline_indexes_markdown_and_docx_end_to_end(client):
    import docx

    _, tok = await signup(client)
    md = (await client.post("/api/knowledge/documents",
                            files={"file": ("policy.md", b"# Remote work\n\nWe work remotely three days a week.",
                                            "text/markdown")},
                            data={"kind": "file"}, headers=auth(tok))).json()
    buf = io.BytesIO()
    d = docx.Document()
    d.add_paragraph("The office address is in Seoul.")
    d.save(buf)
    dx = (await client.post("/api/knowledge/documents",
                            files={"file": ("office.docx", buf.getvalue(),
                                            "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
                            data={"kind": "file"}, headers=auth(tok))).json()
    await _run_jobs(["knowledge.index"])
    for doc_id in (md["id"], dx["id"]):
        got = (await client.get(f"/api/knowledge/documents/{doc_id}", headers=auth(tok))).json()
        assert got["status"] == "ready" and got["chunk_count"] >= 1, got


# ── keys / bootstrap ─────────────────────────────────────────────────────────


async def test_empty_encryption_key_keeps_the_legacy_derivation(client, monkeypatch):
    """Production .env has MEMORA_ENCRYPTION_KEY empty: existing ciphertext must still open."""
    from memora.core import security as SEC
    from memora.services.keycheck import verify_encryption_keys

    settings = SimpleNamespace(encryption_key="", encryption_key_previous="", secret_key="s" * 64,
                               public_url="https://secretary.example")
    monkeypatch.setattr(SEC, "get_settings", lambda: settings)
    SEC.validate_security_settings()  # empty encryption key is a supported production config
    assert SEC.decrypt(SEC.encrypt("stored-secret")) == "stored-secret"
    # and the same bytes still open with the key derived explicitly from the signing key
    assert SEC.legacy_fernet_key(secret_key="s" * 64) == SEC._fernet_key_strings()[0]
    monkeypatch.undo()

    async with session_scope() as db:
        assert await verify_encryption_keys(db) >= 0


async def test_signing_key_rotation_without_pinned_fernet_key_fails_at_startup(client, monkeypatch):
    from memora.core import security as SEC
    from memora.services import settings as S
    from memora.services.keycheck import verify_encryption_keys

    async with session_scope() as db:
        await S.put(db, "providers.anthropic.api_key", "sk-ant-canary")

    rotated = SimpleNamespace(encryption_key="", encryption_key_previous="", secret_key="rotated-" + "z" * 60,
                              public_url="https://secretary.example")
    monkeypatch.setattr(SEC, "get_settings", lambda: rotated)
    async with session_scope() as db:
        with pytest.raises(RuntimeError) as e:
            await verify_encryption_keys(db)
    message = str(e.value)
    assert "MEMORA_ENCRYPTION_KEY" in message
    assert "keytool legacy" in message and "MEMORA_SECRET_KEY" in message

    # pinning the legacy key derived from the OLD signing key repairs it
    from memora.config import get_settings

    rotated.encryption_key_previous = SEC.legacy_fernet_key(secret_key=get_settings().secret_key)
    async with session_scope() as db:
        assert await verify_encryption_keys(db) >= 1
    monkeypatch.undo()
    async with session_scope() as db:
        await S.put(db, "providers.anthropic.api_key", "")
    S.invalidate("providers.")


def test_bootstrap_token_is_only_required_for_https_installs(monkeypatch):
    from memora.services import accounts as A

    monkeypatch.setattr(A, "get_settings",
                        lambda: SimpleNamespace(public_url="http://localhost:3000", bootstrap_token=""))
    A._validate_bootstrap(0, None)  # zero-config local/HTTP bootstrap keeps working

    monkeypatch.setattr(A, "get_settings",
                        lambda: SimpleNamespace(public_url="https://secretary.example", bootstrap_token=""))
    from memora.core.errors import Forbidden

    with pytest.raises(Forbidden) as e:
        A._validate_bootstrap(0, "anything")
    assert e.value.code == "bootstrap_not_configured"
    A._validate_bootstrap(3, None)  # an already-bootstrapped install is unaffected


async def test_unsubscribe_token_is_scoped_to_unsubscribe(client):
    from memora.core.security import sign_state
    from memora.services import notifications as NT

    _, tok = await signup(client)
    channels = (await client.get("/api/notifications/channels", headers=auth(tok))).json()["items"]
    ch_id = channels[0]["id"]
    other = sign_state({"ch": ch_id, "kind": "file_share"})
    r = await client.get("/api/notifications/unsubscribe", params={"token": other})
    assert r.status_code == 403
    good = NT.unsubscribe_token(uuid.UUID(ch_id))
    assert (await client.get("/api/notifications/unsubscribe", params={"token": good})).status_code == 200
    after = (await client.get("/api/notifications/channels", headers=auth(tok))).json()["items"]
    assert [c for c in after if c["id"] == ch_id][0]["enabled"] is False


# ── admin / account-deletion interaction with holds ──────────────────────────


async def test_admin_negative_adjustment_is_refused_cleanly_while_credits_are_held(client):
    from tests.test_audit import _admin_token

    user, _ = await signup(client)
    owner_id = uuid.UUID(user["id"])
    await _fund(owner_id, "6")
    async with session_scope() as db:
        await CR.reserve_turn(db, owner_id=owner_id, turn_id=uuid.uuid4(), turn_cap=6, daily_cap=None,
                              audience="owner")
    admin = await _admin_token(client)
    r = await client.post(f"/api/admin/users/{owner_id}/credits", json={"delta": -5, "note": "clawback"},
                          headers=auth(admin))
    assert r.status_code == 402
    body = r.json()["error"]
    assert body["code"] == "credits_reserved" and body["detail"]["reserved"] == 6.0
    # a positive grant is never blocked by a hold
    r = await client.post(f"/api/admin/users/{owner_id}/credits", json={"delta": 4, "note": "topup"},
                          headers=auth(admin))
    assert r.status_code == 200 and r.json()["balance"] == 10.0


async def test_deleting_an_account_with_a_held_reservation_leaves_no_orphan_rows(client):
    from memora.models import CreditBalance
    from tests.test_audit import _admin_token

    admin = await _admin_token(client)  # first, so the victim below is never the acting admin
    user, _ = await signup(client)
    owner_id = uuid.UUID(user["id"])
    await _fund(owner_id, "8")
    turn_id = uuid.uuid4()
    async with session_scope() as db:
        await CR.reserve_turn(db, owner_id=owner_id, turn_id=turn_id, turn_cap=8, daily_cap=None, audience="owner")

    r = await client.delete(f"/api/admin/users/{owner_id}", headers=auth(admin))
    assert r.status_code == 200, r.text
    async with session_scope() as db:
        assert await db.get(User, owner_id) is None
        assert await db.get(CreditBalance, owner_id) is None
        assert await db.get(CreditReservation, turn_id) is None
        assert await db.scalar(select(func.count()).select_from(CreditLedger)
                               .where(CreditLedger.owner_id == owner_id)) == 0


# ── DNS answer families ──────────────────────────────────────────────────────


def test_safe_http_accepts_public_mixed_and_ipv6_only_answers(monkeypatch):
    import socket as _socket

    from memora.services import safe_http as SH

    monkeypatch.setattr(SH, "get_settings", lambda: SimpleNamespace(outbound_allowed_ports="80,443"))
    monkeypatch.setattr(SH.socket, "getaddrinfo", lambda *_a, **_k: [
        (_socket.AF_INET6, _socket.SOCK_STREAM, 6, "", ("2606:2800:220:1:248:1893:25c8:1946", 443, 0, 0)),
        (_socket.AF_INET, _socket.SOCK_STREAM, 6, "", ("93.184.216.34", 443)),
        (_socket.AF_INET, _socket.SOCK_STREAM, 6, "", ("93.184.216.34", 443)),  # duplicate answer
    ])
    target = SH.resolve_target("https://example.com/")
    assert target.ips == ("2606:2800:220:1:248:1893:25c8:1946", "93.184.216.34")

    monkeypatch.setattr(SH.socket, "getaddrinfo", lambda *_a, **_k: [
        (_socket.AF_INET6, _socket.SOCK_STREAM, 6, "", ("2606:2800:220:1:248:1893:25c8:1946", 443, 0, 0)),
    ])
    assert SH.resolve_target("https://example.com/").ips == ("2606:2800:220:1:248:1893:25c8:1946",)

    # one private answer in a mixed set poisons the whole hostname
    monkeypatch.setattr(SH.socket, "getaddrinfo", lambda *_a, **_k: [
        (_socket.AF_INET, _socket.SOCK_STREAM, 6, "", ("93.184.216.34", 443)),
        (_socket.AF_INET6, _socket.SOCK_STREAM, 6, "", ("::1", 443, 0, 0)),
    ])
    with pytest.raises(ValueError, match="blocked_host"):
        SH.resolve_target("https://example.com/")


async def test_auth_status_separates_bootstrap_needed_from_token_required(client, monkeypatch):
    """The signup form must not demand a token on a zero-config http install."""
    from memora.api import auth as AUTH

    body = (await client.get("/api/auth/status")).json()
    assert body["bootstrap_needed"] is False  # this suite already has accounts
    assert body["bootstrap_token_required"] is False

    monkeypatch.setattr(AUTH.A, "user_count", lambda _db: _zero())
    monkeypatch.setattr(AUTH, "get_settings",
                        lambda: SimpleNamespace(public_url="http://localhost:3000"))
    body = (await client.get("/api/auth/status")).json()
    assert body["bootstrap_needed"] is True and body["bootstrap_token_required"] is False

    monkeypatch.setattr(AUTH, "get_settings",
                        lambda: SimpleNamespace(public_url="https://secretary.example"))
    body = (await client.get("/api/auth/status")).json()
    assert body["bootstrap_needed"] is True and body["bootstrap_token_required"] is True


async def _zero() -> int:
    return 0
